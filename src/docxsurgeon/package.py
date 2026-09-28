"""docxsurgeon：对 OOXML 文件（docx/xlsx）做 zip 包级"手术"。

定位：python-docx 管"文档模型"，docxtpl 管"模板填充"，
本包管"已有文档的包级修改"——替换/编辑嵌入文件、更新图表
数据引用范围、修复复选框打印符号。全部在内存中完成，
不落地临时目录，未修改的部件字节级原样保留。

核心逻辑提取自一套在广电播控机房生产环境运行多年的报表工具，
经过真实 Word/WPS 文档的长期验证。
"""
from __future__ import annotations

import io
import re
import zipfile
from contextlib import contextmanager

__version__ = "0.1.0"

EMBEDDINGS_DIR = "word/embeddings/"


class DocxPackage:
    """把一个 docx（或任意 OOXML zip 包）加载进内存做部件级修改。

    用法::

        pkg = DocxPackage("周报模板.docx")
        with pkg.open_embedded_xlsx("word/embeddings/Microsoft_Excel____.xlsx") as wb:
            ws = wb.active
            ws["A2"] = "9月"
            set_table_range(ws, "A1:C8")
        pkg.set_chart_data_rows("word/charts/chart2.xml", last_row=8)
        pkg.fix_checkbox_glyph()
        pkg.save("输出.docx")
    """

    def __init__(self, source):
        """source 可以是文件路径或已打开的二进制文件对象。"""
        self._order = []   # 保持 zip 内部件的原始顺序
        self._parts = {}   # arcname -> bytes
        with zipfile.ZipFile(source) as zf:
            for info in zf.infolist():
                if info.is_dir():
                    continue
                self._order.append(info.filename)
                self._parts[info.filename] = zf.read(info.filename)

    # ------------------------------------------------------------------
    # 部件（zip 条目）级操作
    # ------------------------------------------------------------------

    def list_parts(self):
        """返回包内全部部件名（按原始 zip 顺序）。"""
        return list(self._order)

    def has_part(self, name):
        return name in self._parts

    def read_part(self, name):
        """按部件名读取原始字节。"""
        self._require_part(name)
        return self._parts[name]

    def replace_part(self, name, data, create=False):
        """整体替换一个部件的内容。data 为 bytes 或 str（按 UTF-8 编码）。

        create=True 时允许写入原包中不存在的新部件（追加到包尾）。
        """
        if not isinstance(data, (bytes, bytearray)):
            data = str(data).encode("utf-8")
        if name not in self._parts:
            if not create:
                raise KeyError("包中不存在部件 %r（如需新增请用 create=True）" % name)
            self._order.append(name)
        self._parts[name] = bytes(data)

    def replace_text_in_part(self, name, old, new):
        """把一个 XML 部件当作文本做子串替换，返回替换次数。"""
        text = self.read_part(name).decode("utf-8")
        count = text.count(old)
        if count:
            self.replace_part(name, text.replace(old, new))
        return count

    # ------------------------------------------------------------------
    # 嵌入文件（word/embeddings/）
    # ------------------------------------------------------------------

    def list_embeddings(self):
        """列出包内全部嵌入文件（word/embeddings/ 下的部件）。"""
        return [p for p in self._order if p.startswith(EMBEDDINGS_DIR)]

    def replace_embedded(self, name, source):
        """用本地文件或字节流替换包内一个嵌入文件。

        name 为 zip 内路径（如 'word/embeddings/Microsoft_Excel____.xlsx'），
        source 为文件路径或 bytes。
        """
        if name not in self._parts:
            raise KeyError("包中不存在嵌入文件 %r，现有：%r" % (name, self.list_embeddings()))
        if isinstance(source, (bytes, bytearray)):
            data = bytes(source)
        else:
            with open(source, "rb") as f:
                data = f.read()
        self._parts[name] = data

    @contextmanager
    def open_embedded_xlsx(self, name):
        """上下文管理器：直接编辑包内嵌入的 xlsx，退出时自动写回。

        yield 一个 openpyxl Workbook；with 块正常结束才写回，
        抛异常则放弃修改（包保持原样）。

        例::

            with pkg.open_embedded_xlsx("word/embeddings/data.xlsx") as wb:
                ws = wb.active
                ws["B2"] = 3.14
        """
        if name not in self._parts:
            raise KeyError("包中不存在嵌入文件 %r，现有：%r" % (name, self.list_embeddings()))
        from openpyxl import load_workbook  # 惰性导入：不用本功能就不需要 openpyxl

        wb = load_workbook(io.BytesIO(self._parts[name]))
        try:
            yield wb
        except Exception:
            raise  # 异常时不写回，避免半成品数据进包
        else:
            buf = io.BytesIO()
            wb.save(buf)
            self._parts[name] = buf.getvalue()

    # ------------------------------------------------------------------
    # 图表数据引用范围
    # ------------------------------------------------------------------

    def set_chart_data_rows(self, chart_part, last_row, sheet="Sheet1", start_row=2):
        """更新图表 XML 中数据系列的行引用范围（$X$start:$X$last_row）。

        对应生产场景"图表引用的行数随数据条数变化"：把 chart XML 里
        所有形如 ``Sheet1!$B$2:$B$20`` / ``Sheet1!$B$2`` 的引用改写为
        以 last_row 结尾（或 last_row <= start_row 时收缩为单行引用）。
        每列字母保持不变，支持多字母列（$AB$）。

        返回改写的引用个数。
        """
        pattern = re.compile(
            r"{sheet}!\$([A-Z]{{1,3}})\${row}(?::\$([A-Z]{{1,3}})\$(\d+))?".format(
                sheet=re.escape(sheet), row=start_row)
        )

        def _repl(m):
            col = m.group(1)
            if last_row <= start_row:
                return "{s}!${c}${r}".format(s=sheet, c=col, r=start_row)
            return "{s}!${c}${r}:${c}${e}".format(s=sheet, c=col, r=start_row, e=last_row)

        text = self.read_part(chart_part).decode("utf-8")
        new_text, count = pattern.subn(_repl, text)
        if count:
            self.replace_part(chart_part, new_text)
        return count

    # ------------------------------------------------------------------
    # 复选框打印符号修复
    # ------------------------------------------------------------------

    def fix_checkbox_glyph(self, part="word/document.xml", search="☑",
                            symbol_font="Wingdings 2", symbol_char="0052",
                            run_font="黑体", cs_font="等线"):
        """把 Unicode 勾选字符替换为 Wingdings 2 符号 run，保证打印正常。

        背景：文档里的 "☑" 字符在部分打印流程中会显示为方框；Word 正确的
        表达方式是一个 ``<w:sym w:font="Wingdings 2" w:char="0052"/>`` 符号。
        本方法把 search 字符从文本 run 中切出来，替换为独立的符号 run。
        返回替换次数。
        """
        replacement = (
            '</w:t></w:r>'
            '<w:r><w:rPr><w:rFonts w:hint="eastAsia" w:ascii="{rf}" '
            'w:hAnsi="{rf}" w:eastAsia="{rf}" w:cs="{cs}"/></w:rPr>'
            '<w:sym w:font="{sf}" w:char="{sc}"/></w:r>'
            '<w:r><w:rPr><w:rFonts w:hint="eastAsia" w:ascii="{rf}" '
            'w:hAnsi="{rf}" w:eastAsia="{rf}" w:cs="{cs}"/></w:rPr>'
            '<w:t xml:space="preserve">'
        ).format(rf=run_font, cs=cs_font, sf=symbol_font, sc=symbol_char)
        return self.replace_text_in_part(part, search, replacement)

    # ------------------------------------------------------------------
    # 保存
    # ------------------------------------------------------------------

    def save(self, path):
        """写回磁盘。未修改的部件按原始顺序、原始字节写入（ZIP_DEFLATED）。"""
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
            for name in self._order:
                zf.writestr(name, self._parts[name])

    # ------------------------------------------------------------------
    # 内部
    # ------------------------------------------------------------------

    def _require_part(self, name):
        if name not in self._parts:
            raise KeyError("包中不存在部件 %r" % name)


def set_table_range(ws, ref):
    """修改工作表中 Excel 表格（ListObject）的数据范围。

    在 open_embedded_xlsx 的 with 块内配合使用::

        with pkg.open_embedded_xlsx(...) as wb:
            ws = wb.active
            ws["A2"] = "1月"
            set_table_range(ws, "A1:C8")

    会更新该工作表上的全部表格，返回表格数量。
    """
    tables = list(getattr(ws, "tables", {}).values())
    if not tables:
        raise ValueError("工作表 %r 中没有 Excel 表格（ListObject），无法设置范围" % ws.title)
    for table in tables:
        table.ref = ref
    return len(tables)
