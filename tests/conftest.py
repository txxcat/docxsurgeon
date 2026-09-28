"""测试样本构造：不依赖真实业务文档，全部动态生成。"""
import io
import zipfile

import pytest
from openpyxl import Workbook
from openpyxl.worksheet.table import Table

DOCUMENT_XML = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
    '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
    '<w:body>'
    '<w:p><w:r><w:t>巡检结果：☑</w:t></w:r></w:p>'
    '<w:p><w:r><w:t>复核结果：☑</w:t></w:r></w:p>'
    '</w:body></w:document>'
)

CHART_XML = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
    '<c:chartSpace xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart">'
    '<c:chart><c:plotArea><c:barChart>'
    '<c:ser>'
    '<c:tx><c:strRef><c:f>Sheet1!$B$1</c:f></c:strRef></c:tx>'
    '<c:cat><c:numRef><c:f>Sheet1!$A$2:$A$20</c:f></c:numRef></c:cat>'
    '<c:val><c:numRef><c:f>Sheet1!$B$2:$B$20</c:f></c:numRef></c:val>'
    '</c:ser>'
    '</c:barChart></c:plotArea></c:chart></c:chartSpace>'
)

CHART_XML_MULTI_COLUMN = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
    '<c:chartSpace xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart">'
    '<c:chart><c:plotArea><c:lineChart><c:ser>'
    '<c:val><c:numRef><c:f>Sheet1!$AB$2:$AB$20</c:f></c:numRef></c:val>'
    '<c:val><c:numRef><c:f>Sheet1!$AC$2:$AC$20</c:f></c:numRef></c:val>'
    '</c:ser></c:lineChart></c:plotArea></c:chart></c:chartSpace>'
)


def make_embedded_xlsx_bytes(sheet="Sheet1", table_ref="A1:B5"):
    """生成一个带 Excel 表格（ListObject）的内存 xlsx。"""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    ws["A1"] = "月份"
    ws["B1"] = "数值"
    for i in range(4):
        ws.cell(row=i + 2, column=1, value="%d月" % (i + 1))
        ws.cell(row=i + 2, column=2, value=(i + 1) * 10)
    ws.add_table(Table(displayName="Table1", ref=table_ref))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture
def docx_path(tmp_path):
    """构造一个含 ☑ 字符、嵌入 xlsx、图表 XML 的最小 docx。"""
    path = tmp_path / "sample.docx"
    with zipfile.ZipFile(path, "w") as zf:
        zf.writestr("[Content_Types].xml", "<Types/>")
        zf.writestr("_rels/.rels", "<Relationships/>")
        zf.writestr("word/document.xml", DOCUMENT_XML)
        zf.writestr("word/embeddings/data.xlsx", make_embedded_xlsx_bytes())
        zf.writestr("word/charts/chart1.xml", CHART_XML)
    return path


@pytest.fixture
def pkg(docx_path):
    from docxsurgeon import DocxPackage
    return DocxPackage(docx_path)
